# -*- coding: utf-8 -*-
# NSKinetics: simulation of Non-Steady state enzyme Kinetics and inhibitory phenomena
# Copyright (C) 2025-, Sarang S. Bhagwat <sarangbhagwat.developer@gmail.com>
#
# This module is under the MIT open-source license. See
# https://github.com/sarangbhagwat/nskinetics/blob/main/LICENSE
# for license details.
"""Offline pure-logic tests of ``analyses/opt_scenarios/build_opt_split12d_workbooks.py``
(no ``load()``, no simulation).

The generator is loaded by file path: its ``nskinetics.biorefineries.isobutanol``
imports live inside ``main()`` and the simulation helpers, and it loads the
stdlib-only ``paths.py`` by file path, so importing it is build-free (probed in
a fresh interpreter). The tests check the OPT_SCENARIOS table, the merge, the
metric-warning split, the workbook writer (on a synthetic template in a temp
dir, needs ``openpyxl``), the reload-verify comparison and the ScenarioSpec
block formatter.

Converted from the script ``analyses/test_build_opt_split12d_workbooks_offline.py``;
its 34 checks are kept, grouped by the function they exercise.
"""
import json
import math
import os

import pytest

from ._support import PKG_DIR, assert_ran_ok, load_by_path, run_isolated, skip_unless_modules

pytestmark = pytest.mark.isobutanol

#: The folder of the opt_* scenario tooling (generator + checker).
OPT_SCENARIOS_DIR = os.path.join(PKG_DIR, 'analyses', 'opt_scenarios')
GENERATOR_PATH = os.path.join(OPT_SCENARIOS_DIR, 'build_opt_split12d_workbooks.py')

OPT_NAMES = ('opt_PI_uninformed', 'opt_PI_TRY_informed', 'opt_IBO_titer',
             'opt_IBO_yield', 'opt_IBO_productivity', 'opt_EtOH_titer',
             'opt_EtOH_yield', 'opt_EtOH_productivity')

#: The scenario-A snapshot / trial values used by the merge and writer checks.
A_SNAPSHOT = {'k_3': 5.81, 'k_7': 1.203, 'k_17': 0.1077, 'K_17': 0.0086}
APPLIED = {'k_3': 1.278, 'k_17': 2.154, 'threshold_conc': 170.9,
           'target_delta': 5.0, 'max_n_spikes': 50}

HEADER = ['Parameter name', 'Element', 'Kind', 'Units', 'Baseline', 'Shape',
          'Lower', 'Midpoint', 'Upper', 'References', 'Load statement']
ROWS = [
    ['Plant annual operating days', 'TEA', 'isolated', 'd', 330, 'Triangular',
     297, 330, 363, None, 'tea.operating_days = x'],
    ['Isobutanol selling unit price (sale)', 'TEA', 'isolated', '$/kg', 1.5,
     'Triangular', 0.9, 1.5, 1.8, 'price note', 'V514.isobutanol_price = x'],
    ['._k_3', 'Kinetics', 'coupled', 'g_per_l_per_h', 5.81, 'triangular',
     4.648, 5.81, 6.972, None, 'V406.nsk_kinetic_model._te.k_3 = x'],
    ['._k_17', 'Kinetics', 'coupled', 'g_per_l_per_h', 44, 'triangular',
     35.2, 44, 52.8, 'k_17 = 44 g/L/h: old note',
     'V406.nsk_kinetic_model._te.k_17 = x'],
]

_SENTINEL = 'NSK_MODULES '


@pytest.fixture(scope='module')
def gen():
    """The generator module, loaded by file path."""
    return load_by_path('build_opt_split12d_workbooks', GENERATOR_PATH)


def test_generator_import_is_build_free():
    """Importing the generator does not import the biorefinery
    (nskinetics.biorefineries.isobutanol.system / upstream biorefineries), and
    imports no nskinetics / biosteam / thermosteam module at all. Probed in a
    fresh interpreter: this pytest process has already imported nskinetics."""
    probe = (
        "import importlib.util, json, sys\n"
        "before = set(sys.modules)\n"
        "spec = importlib.util.spec_from_file_location(\n"
        f"    'build_opt_split12d_workbooks', {GENERATOR_PATH!r})\n"
        "m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)\n"
        f"print({_SENTINEL!r} + json.dumps({{'all': sorted(sys.modules),\n"
        "                                  'new': sorted(set(sys.modules) - before)}))\n")
    proc = run_isolated(['-c', probe], timeout=120, simulation=False)
    out = assert_ran_ok(proc, _SENTINEL)
    line = [ln for ln in out.splitlines() if ln.startswith(_SENTINEL)][-1]
    mods = json.loads(line[len(_SENTINEL):])
    modules, modules_by_generator = mods['all'], mods['new']
    assert ('nskinetics.biorefineries.isobutanol.system' not in modules
            and not any(m == 'biorefineries' or m.startswith('biorefineries.')
                        for m in modules)), modules
    assert not any(m.split('.')[0] in ('nskinetics', 'biorefineries', 'biosteam',
                                       'thermosteam')
                   for m in modules_by_generator), modules_by_generator


def test_generator_resolves_package_paths(gen):
    """The generator resolves the package workbooks and the results root."""
    assert os.path.isfile(gen.B_WORKBOOK), gen.B_WORKBOOK
    assert os.path.normpath(gen.WORKBOOK_DIR) == os.path.normpath(
        os.path.join(PKG_DIR, 'data', 'parameter_distributions'))
    assert os.path.basename(gen.REPRODUCTION_DIR) == 'opt_scenario_reproductions'
    assert (os.path.normpath(os.path.dirname(gen.REPRODUCTION_DIR))
            == os.path.normpath(gen.paths.results_dir()))


def test_opt_scenarios_table(gen):
    """OPT_SCENARIOS lists exactly the eight opt_* scenarios, each a 12d
    A-anchored burden campaign; the two PI scenarios pin PI (seven _rs350
    campaigns + the relay); each non-PI scenario pins its own campaign
    objective; the anchor is scenario A."""
    assert set(gen.OPT_SCENARIOS) == set(OPT_NAMES)
    assert all('ethanol_isobutanol_metabolic_split_12d' in study
               and '_aA_' in study and study.endswith('_burden')
               and isinstance(trial, int)
               for _o, study, trial in gen.OPT_SCENARIOS.values())
    assert gen.OPT_SCENARIOS['opt_PI_uninformed'][0] == 'PI'
    assert gen.OPT_SCENARIOS['opt_PI_TRY_informed'][0] == 'PI'
    assert sum('_rs350_' in study
               for _o, study, _t in gen.OPT_SCENARIOS.values()) == 7
    assert '_rl15c111dc_' in gen.OPT_SCENARIOS['opt_PI_TRY_informed'][1]
    assert all(objective.lower().replace(' ', '_') in study
               for name, (objective, study, _t) in gen.OPT_SCENARIOS.items()
               if not name.startswith('opt_PI'))
    assert gen.ANCHOR_SCENARIO == 'A'


def test_workbook_filename(gen):
    """workbook_filename follows the registry naming."""
    assert (gen.workbook_filename('opt_IBO_titer')
            == 'parameter-distributions_corn_IBO_EtOH_opt_IBO_titer.xlsx')


def test_split_metric_warnings(gen):
    """Fermentation-level warnings block (incl. tracked-only tau);
    system-level (TEA) warnings are tolerated."""
    levels = {'IBO titer': 'kinetic', 'IRR': 'system', 'PI': 'system'}
    blocking, tolerated = gen.split_metric_warnings(
        ['IBO titer', 'IRR', 'tau', 'PI'], levels)
    assert blocking == ['IBO titer', 'tau']
    assert tolerated == ['IRR', 'PI']


def test_merge_applied(gen):
    """merge keeps exactly the snapshot keys (feeding keys dropped); sampled
    names take the trial value; non-sampled names (k_7, K_*) keep the
    snapshot value; every value is a Python float; an unknown non-feeding
    name raises ValueError."""
    merged = gen.merge_applied(A_SNAPSHOT, APPLIED)
    assert set(merged) == set(A_SNAPSHOT)
    assert merged['k_3'] == 1.278 and merged['k_17'] == 2.154
    assert merged['k_7'] == 1.203 and merged['K_17'] == 0.0086
    assert all(type(v) is float for v in merged.values())
    with pytest.raises(ValueError):
        gen.merge_applied(A_SNAPSHOT, {'k_99': 1.0})


@skip_unless_modules('openpyxl')
def test_write_workbook(gen, tmp_path):
    """write_workbook on a synthetic template: kinetic-row count; kinetic
    Baseline <- full_applied value, Triangular +/-20 %, References prefixed
    with the provenance; price row Baseline <- ibo_price with the relative
    spread preserved; non-kinetic rows and the template untouched; a kinetic
    row missing from full_applied raises ValueError and writes nothing."""
    from openpyxl import Workbook, load_workbook

    def _rows(path):
        ws = load_workbook(path, data_only=True).active
        return {r[10]: r for r in ws.iter_rows(min_row=2, values_only=True)}

    merged = gen.merge_applied(A_SNAPSHOT, APPLIED)
    td = str(tmp_path)
    template = os.path.join(td, 'template.xlsx')
    wb = Workbook()
    ws = wb.active
    ws.append(HEADER)
    for row in ROWS:
        ws.append(row)
    wb.save(template)

    out = os.path.join(td, 'out.xlsx')
    n_kin = gen.write_workbook(template, out, merged, 2.0,
                               provenance='trial 1602')
    rows = _rows(out)
    k3 = rows['V406.nsk_kinetic_model._te.k_3 = x']
    k17 = rows['V406.nsk_kinetic_model._te.k_17 = x']
    price = rows['V514.isobutanol_price = x']
    days = rows['tea.operating_days = x']
    assert n_kin == 2
    assert k3[4] == 1.278 and k17[4] == 2.154
    assert (k17[5] == 'Triangular'
            and math.isclose(k17[6], 0.8*2.154) and math.isclose(k17[7], 2.154)
            and math.isclose(k17[8], 1.2*2.154))
    assert (k17[9] == 'trial 1602 | template note: k_17 = 44 g/L/h: old note'
            and k3[9] == 'trial 1602')
    assert (price[4] == 2.0 and math.isclose(price[6], 0.9*2.0/1.5)
            and math.isclose(price[7], 2.0) and math.isclose(price[8], 1.8*2.0/1.5)
            and price[9] == 'price note')
    assert list(days) == ROWS[0]
    assert _rows(template)['V406.nsk_kinetic_model._te.k_17 = x'][4] == 44

    out2 = os.path.join(td, 'out2.xlsx')
    incomplete = {k: v for k, v in merged.items() if k != 'k_17'}
    with pytest.raises(ValueError):
        gen.write_workbook(template, out2, incomplete, 2.0)
    assert not os.path.exists(out2)


def test_assert_reload_matches(gen):
    """Reload deltas are returned per key, all within tol; a delta beyond tol,
    a finite-vs-NaN pair, a key absent from the reload and opposite
    infinities raise RuntimeError; NaN/NaN and same-infinity pairs match
    (delta NaN)."""
    reproduced = dict(ethanol=0.5, isobutanol=1.2, objective=0.2729)
    close = dict(ethanol=0.5002, isobutanol=1.2003, objective=0.27295)
    deltas = gen.assert_reload_matches(reproduced, close, 1e-3)
    assert (set(deltas) == {'ethanol', 'isobutanol', 'objective'}
            and all(0.0 < d < 1e-3 for d in deltas.values()))
    far = dict(close, isobutanol=1.25)
    with pytest.raises(RuntimeError):
        gen.assert_reload_matches(reproduced, far, 1e-3)
    with pytest.raises(RuntimeError):
        gen.assert_reload_matches(reproduced, dict(close, ethanol=math.nan), 1e-3)
    with pytest.raises(RuntimeError):
        gen.assert_reload_matches(reproduced, {'ethanol': 0.5, 'isobutanol': 1.2}, 1e-3)
    nonfinite = gen.assert_reload_matches(
        dict(isobutanol=math.nan, IRR=-math.inf),
        dict(isobutanol=math.nan, IRR=-math.inf), 1e-3)
    assert all(math.isnan(d) for d in nonfinite.values())
    with pytest.raises(RuntimeError):
        gen.assert_reload_matches(dict(IRR=-math.inf), dict(IRR=math.inf), 1e-3)


def test_comparable_values(gen):
    """comparable_values: MPSPs + metrics, convergence diagnostics dropped."""
    values = gen.comparable_values(
        {'ethanol': 0.46, 'isobutanol': math.nan},
        {'IBO titer': 39.6, 'PI': 0.86, 'spike_feed_residual': 1e-5,
         'n_sims_run': 3, 'final_drift': 4e-5})
    assert set(values) == {'ethanol', 'isobutanol', 'IBO titer', 'PI'}
    assert all(type(v) is float for v in values.values())


def test_format_scenario_spec(gen):
    """The ScenarioSpec block evaluates to the expected entry and never sets
    spike_conc / stage_1_max_x (they stay None)."""
    block = gen.format_scenario_spec(
        'opt_PI_TRY_informed',
        'parameter-distributions_corn_IBO_EtOH_opt_PI_TRY_informed.xlsx',
        dict(threshold=189.19244688495075, target=194.19244688495075,
             spike=600.0, max_n_spikes=1),
        dict(ethanol=0.5, isobutanol=1.2), 'PI', 0.8559)
    registry = eval('{' + block + '}', {'ScenarioSpec': dict})
    entry = registry['opt_PI_TRY_informed']
    assert entry == dict(
        name='opt_PI_TRY_informed',
        workbook=('parameter-distributions_corn_IBO_EtOH_'
                  'opt_PI_TRY_informed.xlsx'),
        max_n_spikes=1, threshold_conc=189.19244688495075,
        target_conc=194.19244688495075,
        expected={'ethanol': 0.5, 'isobutanol': 1.2},
        objective_name='PI', objective_value=0.8559)
    assert 'spike_conc' not in block and 'stage_1_max_x' not in block
